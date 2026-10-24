from io import BytesIO

from docx import Document as BuildDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from xml.etree import ElementTree as ET

from backend.app.models.contracts import Document
from backend.app.parser.format_summary import build_format_summary
from backend.app.parser.structure import annotate_structure
from backend.app.parser.styles import parse_styles, resolve_style
from backend.app.parser.word_parser import parse_docx
from backend.app.rules.contracts import CheckRule
from backend.app.rules.detectors import detect_font, detect_size
from backend.app.rules.targets import paragraph_targets


def _paragraph(paragraph_id, text, *, style="Normal", part="document", heading=None, role=None, structure=None, size_pt=12, drawings=None, alignment=None):
    return {
        "paragraph_id": paragraph_id,
        "text": text,
        "style": {"name": style},
        "location": {"part": part},
        "heading": {"level": heading},
        "structure": structure,
        "structure_role": role,
        "format": {"alignment": alignment} if alignment else {},
        "runs": [{"text": text, "font": {"effective": "宋体"}, "size_pt": size_pt}],
        "fields": [],
        "drawings": drawings or [],
    }


def test_resolve_style_keeps_east_asia_from_defaults_when_style_only_sets_ascii() -> None:
    xml = (
        "<w:styles xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        "<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:eastAsia='宋体'/></w:rPr></w:rPrDefault></w:docDefaults>"
        "<w:style w:styleId='Body'><w:name w:val='正文'/><w:rPr><w:rFonts w:ascii='Times New Roman'/></w:rPr></w:style>"
        "</w:styles>"
    )
    parsed = parse_styles(ET.fromstring(xml))
    font = resolve_style("Body", parsed)["font"]
    assert font["ascii"] == "Times New Roman"
    assert font["eastAsia"] == "宋体"


def test_annotate_structure_skips_toc_entries_when_locating_references() -> None:
    paragraphs = [
        _paragraph("p-0001", "目录"),
        _paragraph("p-0002", "摘要\tI", style="toc 1"),
        _paragraph("p-0003", "参考文献\t77", style="toc 1"),
        _paragraph("p-0004", "致谢\t79", style="toc 1"),
        _paragraph("p-0005", "第一章 绪论", style="Heading 1", heading=1),
        _paragraph("p-0006", "这是正文段落。"),
        _paragraph("p-0007", "参考文献", style="毕设-正文-大标题"),
        _paragraph("p-0008", "[1] 张三. 测试文献."),
        _paragraph("p-0009", "致谢"),
    ]
    annotate_structure(paragraphs)
    by_id = {item["paragraph_id"]: item for item in paragraphs}
    assert by_id["p-0003"]["structure"] == "toc"
    assert by_id["p-0004"]["structure"] == "toc"
    assert by_id["p-0006"]["structure"] == "body"
    assert by_id["p-0007"]["structure"] == "references"
    assert by_id["p-0007"]["structure_role"] == "title"
    assert by_id["p-0008"]["structure"] == "references"
    assert by_id["p-0008"].get("structure_role") == "body"


def test_paragraph_targets_do_not_assign_body_or_references_to_table_cells() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "技术类别", part="table", style="毕设-表格内容", size_pt=10.5, structure="table"),
            _paragraph("p-0002", "这是正文段落。", structure="body"),
        ],
    )
    assigned = paragraph_targets(document)
    assert "body" not in assigned["p-0001"]
    assert "references-entry" not in assigned["p-0001"]
    assert "body" in assigned["p-0002"]
    body_size = CheckRule(id="body-size", type="size", target="body", expected={"size": "小四"})
    assert detect_size(document, body_size) == []


def test_toc_style_title_does_not_switch_following_body_into_references() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "参考文献", style="toc 1", structure="toc"),
            _paragraph("p-0002", "这是正文段落。", structure="body"),
        ],
    )
    assigned = paragraph_targets(document)
    assert assigned["p-0001"] == {"toc-body"}
    assert "references-entry" not in assigned["p-0001"]
    assert assigned["p-0002"] == {"body"}


def test_detect_font_uses_east_asia_for_chinese_text_even_if_effective_is_ascii() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[{
            "paragraph_id": "p-0001",
            "text": "[1] 张三. 测试文献.",
            "style": {"name": "Normal"},
            "heading": {"level": None},
            "structure": "references",
            "structure_role": "body",
            "location": {"part": "document"},
            "format": {},
            "runs": [{
                "text": "[1] 张三. 测试文献.",
                "font": {
                    "effective": "Times New Roman",
                    "ascii": "Times New Roman",
                    "east_asia": "宋体",
                },
                "size_pt": 10.5,
            }],
        }],
    )
    rule = CheckRule(id="references-font", type="font", target="references-entry", expected={"font": "宋体"})
    assert detect_font(document, rule) == []


def test_format_summary_uses_east_asia_for_chinese_body() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[{
            "paragraph_id": "p-0001",
            "text": "正文段落",
            "structure": "body",
            "heading": {"level": None},
            "format": {"line_spacing": 20},
            "runs": [{
                "text": "正文段落",
                "font": {
                    "effective": "Times New Roman",
                    "ascii": "Times New Roman",
                    "east_asia": "宋体",
                },
                "size_pt": 12,
            }],
        }],
    )
    summary = build_format_summary(document)
    assert summary["structured"]["body"]["font"] == "宋体"
    assert "字体：宋体" in summary["text"]


def _save_docx(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _set_run_fonts(run, *, ascii_name: str, east_asia: str) -> None:
    run.font.name = ascii_name
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), ascii_name)
    fonts.set(qn("w:hAnsi"), ascii_name)
    fonts.set(qn("w:eastAsia"), east_asia)


def test_parse_docx_skips_toc_references_entry_and_reads_east_asia_font() -> None:
    source = BuildDocument()
    try:
        source.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    except ValueError:
        pass
    source.add_paragraph("目录")
    source.add_paragraph("参考文献\t77", style="toc 1")
    source.add_paragraph("致谢\t79", style="toc 1")
    heading = source.add_paragraph("第一章 绪论")
    heading.style = "Heading 1"
    source.add_paragraph("这是正文段落。")
    source.add_paragraph("参考文献")
    ref = source.add_paragraph()
    run = ref.add_run("[1] 张三. 测试文献.")
    _set_run_fonts(run, ascii_name="Times New Roman", east_asia="宋体")
    source.add_paragraph("致谢")
    parsed = parse_docx(BytesIO(_save_docx(source)))
    body = [item for item in parsed.paragraphs if item["location"]["part"] == "document"]
    by_text = {item["text"]: item for item in body}
    assert by_text["参考文献\t77"]["structure"] == "toc"
    assert by_text["这是正文段落。"]["structure"] == "body"
    assert by_text["参考文献"]["structure"] == "references"
    assert by_text["参考文献"]["structure_role"] == "title"
    assert by_text["[1] 张三. 测试文献."]["structure"] == "references"
    font = by_text["[1] 张三. 测试文献."]["runs"][0]["font"]
    assert font["east_asia"] == "宋体"
    assert font["ascii"] == "Times New Roman"
    assert font["effective"] == "宋体"


def test_body_after_toc_title_is_not_assigned_toc_body() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "目录", style="毕设-目录-大标题", structure="body"),
            _paragraph("p-0002", "第1章 绪论\t1", style="toc 1", structure="toc"),
            _paragraph("p-0003", "第1章 绪论", style="毕设-正文-大标题", structure="body"),
            _paragraph("p-0004", "这是正文段落。", style="毕设-正文-正文部分", structure="body"),
            _paragraph("p-0005", "1.1 研究背景", style="毕设-正文-1级标题", structure="body"),
        ],
    )
    assigned = paragraph_targets(document)
    assert "toc-body" not in assigned["p-0003"]
    assert "title1" in assigned["p-0003"]
    assert assigned["p-0004"] == {"body"}
    assert "toc-body" not in assigned["p-0005"]
    assert "title2" in assigned["p-0005"]


def test_font_for_text_uses_east_asia_for_cjk_punctuation() -> None:
    from backend.app.parser.styles import font_for_text

    font = {"ascii": "Times New Roman", "east_asia": "宋体", "effective": "Times New Roman"}
    assert font_for_text(font, "。") == "宋体"
    assert font_for_text(font, "，") == "宋体"
    assert font_for_text(font, "“") == "宋体"


def test_detect_font_allows_times_new_roman_on_latin_run_when_expected_song() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[{
            "paragraph_id": "p-0001",
            "text": "PostgreSQL 数据库",
            "style": {"name": "Normal"},
            "heading": {"level": None},
            "structure": "body",
            "location": {"part": "document"},
            "format": {},
            "runs": [
                {"text": "PostgreSQL", "font": {"effective": "Times New Roman", "ascii": "Times New Roman", "east_asia": "宋体"}, "size_pt": 12},
                {"text": "数据库", "font": {"effective": "Times New Roman", "ascii": "Times New Roman", "east_asia": "宋体"}, "size_pt": 12},
            ],
        }],
    )
    rule = CheckRule(id="body-font", type="font", target="body", expected={"font": "宋体"})
    errors = detect_font(document, rule)
    assert errors == []


def test_page_margin_accepts_word_twip_rounding() -> None:
    from backend.app.rules.detectors import detect_page_margin

    document = Document(document_id="D1", source_filename="x.docx", sections=[{
        "index": 0,
        "margins_pt": {"top": 99.25, "right": 85.05, "bottom": 99.25, "left": 85.05},
        "header_distance_pt": 77.95,
        "footer_distance_pt": 49.6,
        "page_width_pt": 595.3,
        "page_height_pt": 841.9,
    }])
    errors = detect_page_margin(document, CheckRule(id="page-margin", type="page_margin", expected={
        "margin": {"top": "3.5cm", "right": "3.0cm", "bottom": "3.5cm", "left": "3.0cm"},
        "header": "2.75cm",
        "footer": "1.75cm",
        "page": {"width_cm": 21.0, "height_cm": 29.7},
    }))
    assert errors == []


def test_heading_numbering_ignores_toc_chapter_entries() -> None:
    from backend.app.rules.detectors import detect_heading_numbering

    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "第1章 绪论\t1", style="toc 1", structure="toc", heading=1),
            _paragraph("p-0002", "第7章 总结\t70", style="toc 1", structure="toc", heading=1),
            _paragraph("p-0003", "第1章 绪论", style="毕设-正文-大标题", structure="body", heading=1),
        ],
    )
    errors = detect_heading_numbering(document, CheckRule(id="heading-numbering", type="heading_numbering", expected={}))
    assert errors == []


def test_cover_title_prefers_topic_line_over_larger_school_name() -> None:
    paragraphs = [
        _paragraph("p-0001", "电子科技大学成都学院", size_pt=36),
        _paragraph("p-0002", "毕业论文（设计）", size_pt=36),
        _paragraph("p-0003", "题    目 基于Django的调酒配方与", size_pt=18),
        _paragraph("p-0004", "库存一体化平台设计与实现", size_pt=16),
        _paragraph("p-0005", "学    院        计算机学院", size_pt=18),
        _paragraph("p-0006", "摘要"),
        _paragraph("p-0007", "摘要正文"),
    ]
    annotate_structure(paragraphs)
    by_id = {item["paragraph_id"]: item for item in paragraphs}
    assert by_id["p-0001"]["structure"] == "cover"
    assert by_id["p-0001"].get("structure_role") != "title"
    assert by_id["p-0002"].get("structure_role") != "title"
    assert by_id["p-0003"]["structure"] == "cover"
    assert by_id["p-0003"]["structure_role"] == "title"
    assert by_id["p-0005"].get("structure_role") != "title"


def test_image_note_and_empty_drawing_paragraphs_are_not_body_targets() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "", style="毕设-图片", structure="body", size_pt=10.5, drawings=[{"kind": "drawing"}], alignment="center"),
            _paragraph("p-0002", "续表4-5库存原料表", style="毕设-图注", structure="body", size_pt=10.5, alignment="center"),
            _paragraph("p-0003", "", style="No Spacing", structure="body", size_pt=10.5, drawings=[{"kind": "drawing"}], alignment="center"),
            _paragraph("p-0004", "", structure="body", size_pt=11, alignment="center"),
            _paragraph("p-0005", "这是正文段落。", structure="body", size_pt=12, alignment="justify"),
        ],
    )
    assigned = paragraph_targets(document)
    assert "body" not in assigned["p-0001"]
    assert "body" not in assigned["p-0002"]
    assert "table_caption" in assigned["p-0002"] or "table-caption" in assigned["p-0002"]
    assert "body" not in assigned["p-0003"]
    assert "body" not in assigned["p-0004"]
    assert assigned["p-0005"] == {"body"}
    body_size = CheckRule(id="body-size", type="size", target="body", expected={"size": "小四"})
    errors = detect_size(document, body_size)
    assert errors == []


def test_foreign_section_abstract_headings_are_not_rematched() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "摘要", structure="abstract", role="title"),
            _paragraph("p-0002", "摘要正文", structure="abstract", role="body"),
            _paragraph("p-0003", "第1章 绪论", style="毕设-正文-大标题", structure="body", heading=1),
            _paragraph("p-0004", "Abstract", style="毕业@外文|正文", structure="body"),
            _paragraph("p-0005", "摘要", structure="body"),
            _paragraph("p-0006", "外文资料原文", style="毕业@外文|大标题", structure="body"),
        ],
    )
    assigned = paragraph_targets(document)
    assert assigned["p-0001"] == {"abstract-title"}
    assert "abstract-en-title" not in assigned["p-0004"]
    assert "abstract-title" not in assigned["p-0005"]
    assert "body" not in assigned["p-0004"]
    assert "body" not in assigned["p-0006"]
    assert "body" in assigned["p-0005"]


def test_page_number_detector_accepts_footer_after_unnumbered_cover_section() -> None:
    from backend.app.rules.detectors import detect_page_number

    document = Document(document_id="D1", source_filename="x.docx", pages=[
        {"section_index": 0, "page_index": 0, "page_number": 1, "position": "none", "fields": []},
        {"section_index": 1, "page_index": 1, "page_number": 1, "position": "footer", "fields": [{"instruction": "PAGE", "part": "footer"}]},
        {"section_index": 1, "page_index": 2, "page_number": 2, "position": "footer", "fields": [{"instruction": "PAGE", "part": "footer"}]},
    ])
    errors = detect_page_number(document, CheckRule(id="page-number", type="page_number", expected={"position": "footer", "continuous": True}))
    assert errors == []


def test_real_preface_titles_are_kept_before_chapters() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "摘要", structure="abstract", role="title"),
            _paragraph("p-0002", "ABSTRACT", style="毕设-摘要标题", structure="body"),
            _paragraph("p-0003", "目 录", style="毕设-目录-大标题", structure="body"),
            _paragraph("p-0004", "第1章 绪论", style="毕设-正文-大标题", structure="body", heading=1),
            _paragraph("p-0005", "摘要", structure="body"),
        ],
    )
    assigned = paragraph_targets(document)
    assert "abstract-en-title" in assigned["p-0002"]
    assert "toc-title" in assigned["p-0003"]
    assert "abstract-title" not in assigned["p-0005"]
    assert "body" in assigned["p-0005"]


def test_heading_numbering_ignores_chapter_mentions_in_body_text() -> None:
    from backend.app.rules.detectors import detect_heading_numbering

    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "第1章 绪论", style="毕设-正文-大标题", structure="body", heading=1),
            _paragraph("p-0002", "1.1 选题背景", structure="body", heading=2),
            _paragraph("p-0003", "1.2 国内外研究现状", structure="body", heading=2),
            _paragraph("p-0004", "1.3 论文章节安排", structure="body", heading=2),
            _paragraph("p-0005", "第1章为绪论，对线上自习室小程序的研究背景、意义和研究现状进行全面阐述。", structure="body"),
            _paragraph("p-0006", "第2章 系统相关技术背景，介绍小程序开发所用到的关键技术。", structure="body"),
            _paragraph("p-0007", "第2章 系统相关技术背景", style="毕设-正文-大标题", structure="body", heading=1),
            _paragraph("p-0008", "第3章 系统需求分析", style="毕设-正文-大标题", structure="body", heading=1),
        ],
    )
    rule = CheckRule(id="heading-numbering", type="heading_numbering", expected={"require_title": True})
    assert detect_heading_numbering(document, rule) == []


def test_heading_numbering_still_detects_gap_for_unstyled_chapter_titles() -> None:
    from backend.app.rules.detectors import detect_heading_numbering

    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "第1章 绪论", structure="body"),
            _paragraph("p-0002", "第3章 系统需求分析", structure="body"),
        ],
    )
    rule = CheckRule(id="heading-numbering", type="heading_numbering", expected={"require_title": True})
    errors = detect_heading_numbering(document, rule)
    assert [error.current for error in errors] == ["3"]
    assert [error.expected for error in errors] == ["2"]


def test_caption_note_alignment_ignores_empty_paragraph_after_figure_caption() -> None:
    from backend.app.rules.detectors import detect_alignment

    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "图3-4 查看自习室功能结构图", style="新图样式", structure="figure_caption", alignment="center"),
            _paragraph("p-0002", "", style="毕设-图注", structure="body", alignment="justify"),
        ],
    )
    rule = CheckRule(id="caption-note-alignment", type="alignment", target="caption-note", expected={"alignment": "center"})
    assert detect_alignment(document, rule) == []


def test_references_spacing_ignores_empty_paragraph_at_section_end() -> None:
    from backend.app.rules.detectors import detect_paragraph_spacing

    empty = _paragraph("p-0003", "", style="毕设-正文-正文部分", structure="references", role="body", alignment="justify")
    empty["format"]["space_before_pt"] = 6.0
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "参考文献", style="毕设-正文-大标题", structure="references", role="title", alignment="center"),
            _paragraph("p-0002", "张三. 测试文献[J]. 2020.", style="参考文献", structure="references", role="body", alignment="justify"),
            empty,
        ],
    )
    rule = CheckRule(id="references-spacing", type="paragraph_spacing", target="references-entry",
                     expected={"space_before_pt": 3, "space_after_pt": 0})
    assert detect_paragraph_spacing(document, rule) == []


def test_references_spacing_still_checks_real_entries() -> None:
    from backend.app.rules.detectors import detect_paragraph_spacing

    entry = _paragraph("p-0002", "张三. 测试文献[J]. 2020.", style="参考文献", structure="references", role="body", alignment="justify")
    entry["format"]["space_before_pt"] = 6.0
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[
            _paragraph("p-0001", "参考文献", style="毕设-正文-大标题", structure="references", role="title", alignment="center"),
            entry,
        ],
    )
    rule = CheckRule(id="references-spacing", type="paragraph_spacing", target="references-entry",
                     expected={"space_before_pt": 3, "space_after_pt": 0})
    errors = detect_paragraph_spacing(document, rule)
    assert [error.current for error in errors] == ["段前 6 磅"]
    assert [error.expected for error in errors] == ["段前 3 磅"]


def test_font_detection_accepts_wps_theme_font_name_with_suffix() -> None:
    from backend.app.rules.detectors import detect_font

    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[{
            "paragraph_id": "p-0001",
            "text": "这是正文段落。",
            "style": {"name": "Normal"},
            "heading": {"level": None},
            "structure": "body",
            "location": {"part": "document"},
            "format": {},
            "runs": [{
                "text": "这是正文段落。",
                "font": {"effective": "宋体（正文）", "ascii": "Times New Roman", "east_asia": "宋体（正文）"},
                "size_pt": 12,
            }],
        }],
    )
    rule = CheckRule(id="body-font", type="font", target="body", expected={"font": "宋体"})
    assert detect_font(document, rule) == []


def test_font_detection_accepts_wps_heading_theme_font_name_with_suffix() -> None:
    from backend.app.rules.detectors import detect_font

    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[{
            "paragraph_id": "p-0001",
            "text": "第1章 绪论",
            "style": {"name": "Heading 1"},
            "heading": {"level": 1},
            "structure": "body",
            "location": {"part": "document"},
            "format": {},
            "runs": [{
                "text": "第1章 绪论",
                "font": {"effective": "黑体（标题）", "ascii": "Times New Roman", "east_asia": "黑体（标题）"},
                "size_pt": 15,
            }],
        }],
    )
    rule = CheckRule(id="title1-font", type="font", target="title1", expected={"font": "黑体"})
    assert detect_font(document, rule) == []


def test_cover_title_ignores_topic_label_run() -> None:
    document = Document(
        document_id="D1",
        source_filename="x.docx",
        paragraphs=[{
            "paragraph_id": "p-0001",
            "text": "题    目 面向格式检测的论文题目",
            "style": {"name": "Normal"},
            "heading": {"level": None},
            "structure": "cover",
            "structure_role": "title",
            "location": {"part": "document"},
            "format": {},
            "runs": [
                {"text": "题    目", "font": {"effective": "黑体", "east_asia": "黑体", "ascii": "黑体"}, "size_pt": 18},
                {"text": "面向格式检测的论文题目", "font": {"effective": "宋体", "east_asia": "宋体", "ascii": "Times New Roman"}, "size_pt": 16},
            ],
        }],
    )
    font_rule = CheckRule(id="cover-title-font", type="font", target="cover_title", expected={"font": "宋体"})
    size_rule = CheckRule(id="cover-title-size", type="size", target="cover_title", expected={"size": "三号"})
    assert detect_font(document, font_rule) == []
    assert detect_size(document, size_rule) == []
