from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt

from backend.app.services.docx_parser import parse_docx


def make_sample_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(3.5)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(3)
    section.header_distance = Cm(2.75)
    section.footer_distance = Cm(1.75)
    section.header.paragraphs[0].text = "学校论文规范"
    section.footer.paragraphs[0].text = "第 1 页"

    heading = document.styles["Heading 1"]
    heading.font.name = "黑体"
    heading.font.size = Pt(16)

    title = document.add_paragraph("第一章 绪论", style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.name = "黑体"
    title.runs[0]._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "黑体",
    )
    title.runs[0].font.size = Pt(16)

    body = document.add_paragraph("这是正文内容。")
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = Pt(24)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    body.paragraph_format.line_spacing = Pt(20)
    run = body.runs[0]
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "宋体",
    )
    run.font.size = Pt(12)
    run.bold = True

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "字体"
    table.cell(1, 1).text = "宋体"
    document.save(path)


def test_parse_docx_extracts_document_structure_and_format(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    make_sample_docx(source)

    parsed = parse_docx(source)

    assert parsed["source_filename"] == "sample.docx"
    assert parsed["sections"][0]["margins_cm"] == {
        "top": 3.5,
        "right": 3.0,
        "bottom": 3.5,
        "left": 3.0,
    }
    assert parsed["sections"][0]["header_distance_cm"] == 2.75
    assert parsed["sections"][0]["footer_distance_cm"] == 1.75
    assert parsed["headers"][0]["text"] == "学校论文规范"
    assert parsed["footers"][0]["text"] == "第 1 页"

    heading = parsed["paragraphs"][0]
    assert heading["text"] == "第一章 绪论"
    assert heading["heading_level"] == 1
    assert heading["alignment"] == "center"
    assert heading["runs"][0]["font"] == "黑体"
    assert heading["runs"][0]["size_pt"] == 16.0

    body = parsed["paragraphs"][1]
    assert body["text"] == "这是正文内容。"
    assert body["alignment"] == "justify"
    assert body["first_line_indent_pt"] == 24.0
    assert body["line_spacing_pt"] == 20.0
    assert body["runs"][0]["font"] == "宋体"
    assert body["runs"][0]["bold"] is True

    assert parsed["tables"] == [
        {"index": 0, "rows": [["项目", "结果"], ["字体", "宋体"]]}
    ]


def test_parse_docx_resolves_formatting_from_paragraph_and_run_styles(tmp_path: Path) -> None:
    source = tmp_path / "style-only.docx"
    document = Document()

    paragraph_style = document.styles.add_style("Configured Body", WD_STYLE_TYPE.PARAGRAPH)
    paragraph_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph_style.paragraph_format.first_line_indent = Pt(24)
    paragraph_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph_style.paragraph_format.line_spacing = Pt(20)
    paragraph_style.paragraph_format.space_before = Pt(6)
    paragraph_style.paragraph_format.space_after = Pt(0)

    character_style = document.styles.add_style("Configured Text", WD_STYLE_TYPE.CHARACTER)
    character_style.font.name = "宋体"
    character_style.font.size = Pt(12)
    character_style.font.bold = False

    paragraph = document.add_paragraph("样式正文", style="Configured Body")
    paragraph.runs[0].style = character_style
    document.save(source)

    parsed = parse_docx(source)

    assert parsed["paragraphs"][0]["alignment"] == "justify"
    assert parsed["paragraphs"][0]["space_before_pt"] == 6.0
    assert parsed["paragraphs"][0]["space_after_pt"] == 0.0
    assert parsed["paragraphs"][0]["first_line_indent_pt"] == 24.0
    assert parsed["paragraphs"][0]["line_spacing_rule"] == "exactly"
    assert parsed["paragraphs"][0]["line_spacing_pt"] == 20.0
    assert parsed["paragraphs"][0]["runs"][0]["font"] == "宋体"
    assert parsed["paragraphs"][0]["runs"][0]["size_pt"] == 12.0
    assert parsed["paragraphs"][0]["runs"][0]["bold"] is False
