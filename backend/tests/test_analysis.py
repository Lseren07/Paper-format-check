from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from fastapi.testclient import TestClient

from backend.app.api import upload
from backend.app.main import app
from backend.app.services.task_store import TaskRecord, store
from backend.app.parser.format_summary import line_spacing_label

BODY_TEXT = "随着人工智能技术的发展，论文格式检测逐渐自动化。"
TITLE_TEXT = "基于深度学习的论文格式检测系统设计"
UPLOAD_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def set_font(run, name: str, points: float) -> None:
    run.font.name = name
    run.font.size = Pt(points)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def make_thesis_docx() -> bytes:
    document = Document()
    title = document.add_paragraph(TITLE_TEXT, style="Heading 1")
    set_font(title.runs[0], "黑体", 16)

    body = document.add_paragraph()
    set_font(body.add_run(BODY_TEXT), "宋体", 12)
    body.paragraph_format.line_spacing = 1.5

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def upload_thesis(client: TestClient) -> str:
    response = client.post(
        "/api/v1/paper/upload",
        files={"file": ("thesis.docx", make_thesis_docx(), UPLOAD_CONTENT_TYPE)},
    )
    assert response.status_code == 200
    return response.json()["data"]["task_id"]


def detect(client: TestClient, task_id: str) -> None:
    started = client.post("/api/v1/detect/start", json={"task_id": task_id})
    assert started.status_code == 200
    assert started.json()["data"]["status"] == "completed"


def test_analysis_unknown_task_returns_404() -> None:
    response = TestClient(app).get("/api/v1/document/analysis/Tmissing")
    assert response.status_code == 404
    assert response.json()["detail"] == "检测任务不存在"


def test_analysis_before_detection_returns_409(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(upload, "UPLOAD_DIR", tmp_path)
    client = TestClient(app)
    task_id = upload_thesis(client)

    response = client.get(f"/api/v1/document/analysis/{task_id}")
    assert response.status_code == 409
    assert response.json()["detail"] == "检测尚未完成"


def test_analysis_failed_task_returns_400() -> None:
    store.put(TaskRecord(task_id="Tfailed1", filename="broken.docx", status="failed", message="文档解析失败"))
    try:
        response = TestClient(app).get("/api/v1/document/analysis/Tfailed1")
    finally:
        store.delete("Tfailed1")
    assert response.status_code == 400
    assert response.json()["detail"] == "文档解析失败"


def test_analysis_returns_format_summary(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(upload, "UPLOAD_DIR", tmp_path)
    client = TestClient(app)
    task_id = upload_thesis(client)
    detect(client, task_id)

    response = client.get(f"/api/v1/document/analysis/{task_id}")
    assert response.status_code == 200
    body = response.json()
    assert body["code"] == 200
    data = body["data"]
    assert data["task_id"] == task_id
    assert data["filename"] == "thesis.docx"
    assert data["status"] == "completed"

    # 正文与标题的字体字号行距按接口文档第 7 章聚合
    assert data["format_summary"]["body"] == {"font": "宋体", "size": "小四", "line_spacing": "1.5倍"}
    assert data["format_summary"]["title"]["line_spacing"] == line_spacing_label(1.15)
    # python-docx 默认模板左右边距与上下不同，逐边输出
    margin = data["format_summary"]["page"]["margin"]
    assert "2.54cm" in margin and "3.17cm" in margin
    assert data["format_text"].startswith("正文格式：")


def test_analysis_never_exposes_paper_text(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(upload, "UPLOAD_DIR", tmp_path)
    client = TestClient(app)
    task_id = upload_thesis(client)
    detect(client, task_id)

    body = client.get(f"/api/v1/document/analysis/{task_id}").json()["data"]
    payload = str(body)
    assert BODY_TEXT not in payload
    assert TITLE_TEXT not in payload
    assert "document" not in body
    assert "paragraphs" not in payload
    assert "runs" not in payload
