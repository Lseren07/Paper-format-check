"""规则集注册表：目录扫描 + 可配置选择，规则文件不再是硬编码的单个文件。"""

import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from backend.app.api import upload
from backend.app.main import app
from backend.app.rules.registry import (
    RuleSetNotFoundError,
    clear_rule_cache,
    default_rule_set_id,
    discover_rule_sets,
    load_rule_set,
    resolve_rule_set,
    rules_dir,
)
from backend.app.services import report as report_service
from backend.app.services.task_store import store

REPO_RULES = Path(__file__).parents[2] / "rules"
# 测试里另备一套规则集，用来验证「选了它就不会用默认的」。
SECOND_ID = "school-b"

# 没有 `checks` 的旧结构（page / body / headings / manual_checks），由加载器翻译。
OLD_SHAPE_RULE = {
    "schema_version": "1.0",
    "name": "旧结构规则集",
    "source": "附件1：电子科技大学成都学院毕业论文（设计）撰写格式规范.docx",
    "page": {
        "width_cm": 21.0,
        "height_cm": 29.7,
        "margins_cm": {"top": 3.5, "right": 3.0, "bottom": 3.5, "left": 3.0},
        "header_distance_cm": 2.75,
        "footer_distance_cm": 1.75,
    },
    "body": {
        "font": "宋体",
        "size_pt": 12.0,
        "alignment": "justify",
        "first_line_indent_pt": 24.0,
        "line_spacing_pt": 20.0,
        "space_before_pt": 6.0,
        "space_after_pt": 0.0,
    },
    "headings": {
        "1": {"font": "黑体", "size_pt": 15.0, "bold": True, "alignment": "center", "space_before_pt": 30.0, "space_after_pt": 30.0},
        "4": {"font": "黑体", "size_pt": 12.0, "alignment": "left", "space_before_pt": 6.0, "space_after_pt": 6.0},
    },
    "manual_checks": ["参考文献著录内容", "图表编号和标题"],
}


@pytest.fixture(autouse=True)
def _isolated_cache():
    clear_rule_cache()
    yield
    clear_rule_cache()


def _rule_payload(name: str, **extra) -> dict:
    payload = {
        "schema_version": "1.0",
        "name": name,
        "checks": [{"id": "body-font", "type": "font", "target": "body", "expected": {"font": "宋体"}}],
    }
    payload.update(extra)
    return payload


def _second_rule_payload() -> dict:
    """一套与 default.json 要求不同的规则：正文要求黑体（默认规则要求宋体）。"""
    return _rule_payload("School B", checks=[
        {"id": "body-font", "type": "font", "target": "body", "expected": {"font": "黑体"}},
    ])


def _write_rule(directory: Path, filename: str, payload: dict) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / filename
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def _prepare_rules_dir(tmp_path: Path, *, second: dict | None = None) -> Path:
    """造一个规则目录：默认规则集的副本 + 可选的第二个规则集。"""
    directory = tmp_path / "rules"
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "default.json").write_bytes((REPO_RULES / "default.json").read_bytes())
    if second is not None:
        _write_rule(directory, f"{SECOND_ID}.json", second)
    return directory


def test_rules_dir_defaults_to_repo_rules(monkeypatch) -> None:
    monkeypatch.delenv("PAPER_RULES_DIR", raising=False)
    monkeypatch.delenv("PAPER_RULES_DEFAULT", raising=False)

    assert rules_dir() == REPO_RULES
    assert default_rule_set_id() == "default"


def test_discover_lists_every_rule_file_in_directory(monkeypatch) -> None:
    monkeypatch.setenv("PAPER_RULES_DIR", str(REPO_RULES))

    ids = {info.id for info in discover_rule_sets()}

    assert ids == {path.stem for path in REPO_RULES.glob("*.json")}
    assert "default" in ids


def test_every_repo_rule_file_is_loadable() -> None:
    """目录扫描的意义在于每个文件都能被真正加载，不能有沉默的坏文件。"""
    infos = discover_rule_sets(REPO_RULES)

    assert infos
    assert [info.id for info in infos if not info.loadable] == []


def test_default_rule_set_is_the_one_shipped_with_repo(monkeypatch) -> None:
    monkeypatch.setenv("PAPER_RULES_DIR", str(REPO_RULES))
    monkeypatch.delenv("PAPER_RULES_DEFAULT", raising=False)

    resolved = resolve_rule_set()

    assert resolved.info.id == "default"
    assert resolved.info.file == "default.json"
    assert resolved.info.enabled_check_count > 1
    assert resolved.rules.source is not None


def test_old_shape_rule_file_is_translated_into_real_checks(tmp_path, monkeypatch) -> None:
    """没有 `checks` 的旧结构文件也能被扫描并翻译成可执行规则。"""
    directory = _prepare_rules_dir(tmp_path)
    _write_rule(directory, "old-shape.json", OLD_SHAPE_RULE)
    monkeypatch.setenv("PAPER_RULES_DIR", str(directory))

    rules = load_rule_set("old-shape")

    kinds = {check.type for check in rules.checks}
    assert {"font", "size", "alignment", "line_spacing", "paragraph_indent", "paragraph_spacing", "page_margin"} <= kinds
    assert {str(check.target) for check in rules.checks} >= {"body", "title1", "title4"}
    expected = {check.id: check.expected for check in rules.checks}
    assert expected["body-first-line-indent"] == {"first_line_indent": "2字符"}
    assert expected["page-margin"]["margin"]["top"] == "3.5cm"
    # manual_checks 落成人工核对提示，而不是凭空生成条款
    assert rules.source is not None
    assert rules.source.clauses == []
    assert "参考文献著录内容" in rules.source.notes


def test_unknown_rule_set_id_raises(monkeypatch) -> None:
    monkeypatch.setenv("PAPER_RULES_DIR", str(REPO_RULES))

    with pytest.raises(RuleSetNotFoundError):
        load_rule_set("no-such-school")


def test_default_rule_set_is_configurable(tmp_path, monkeypatch) -> None:
    _write_rule(tmp_path, "alpha.json", _rule_payload("Alpha"))
    _write_rule(tmp_path, "beta.json", _rule_payload("Beta", rule_set_id="beta"))
    monkeypatch.setenv("PAPER_RULES_DIR", str(tmp_path))
    monkeypatch.setenv("PAPER_RULES_DEFAULT", "beta")

    assert resolve_rule_set().rules.name == "Beta"
    assert resolve_rule_set("alpha").rules.name == "Alpha"


def test_rule_set_id_falls_back_to_filename(tmp_path, monkeypatch) -> None:
    _write_rule(tmp_path, "some-school.json", _rule_payload("Some School"))
    monkeypatch.setenv("PAPER_RULES_DIR", str(tmp_path))
    monkeypatch.delenv("PAPER_RULES_DEFAULT", raising=False)

    info = resolve_rule_set("some-school").info

    assert info.id == "some-school"
    assert info.file == "some-school.json"
    assert info.enabled_check_count == 1


def test_broken_file_is_reported_without_hiding_others(tmp_path, monkeypatch) -> None:
    directory = tmp_path
    (directory / "broken.json").write_text("{not json", encoding="utf-8")
    _write_rule(directory, "good.json", _rule_payload("Good"))
    monkeypatch.setenv("PAPER_RULES_DIR", str(directory))

    infos = {info.id: info for info in discover_rule_sets()}

    assert infos["broken"].loadable is False
    assert infos["broken"].error
    assert infos["good"].loadable is True
    with pytest.raises(RuleSetNotFoundError):
        load_rule_set("broken")


def test_rule_list_endpoint_exposes_catalog(monkeypatch) -> None:
    monkeypatch.setenv("PAPER_RULES_DIR", str(REPO_RULES))
    monkeypatch.delenv("PAPER_RULES_DEFAULT", raising=False)

    body = TestClient(app).get("/api/v1/rule/list").json()

    assert body["code"] == 200
    data = body["data"]
    assert data["default_rule_set_id"] == "default"
    assert data["total"] == len(data["rule_sets"]) >= 1
    by_id = {item["id"]: item for item in data["rule_sets"]}
    assert by_id["default"]["file"] == "default.json"
    assert by_id["default"]["enabled_check_count"] > 0
    assert by_id["default"]["source_available"] is True


def test_rule_detail_endpoint_returns_checks_and_404_for_unknown(monkeypatch) -> None:
    monkeypatch.setenv("PAPER_RULES_DIR", str(REPO_RULES))
    client = TestClient(app)

    ok = client.get("/api/v1/rule/default")
    assert ok.status_code == 200
    assert ok.json()["data"]["rule_set"]["id"] == "default"
    assert ok.json()["data"]["checks"]

    missing = client.get("/api/v1/rule/nope")
    assert missing.status_code == 404


def _upload_sample(client: TestClient) -> str:
    """正文用宋体的小样例：默认规则要求宋体（通过），school-b 要求黑体（报错）。"""
    document = Document()
    document.add_paragraph("第一章 绪论", style="Heading 1")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("正文内容示例")
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    stream = BytesIO()
    document.save(stream)
    response = client.post(
        "/api/v1/paper/upload",
        files={"file": ("thesis.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    return response.json()["data"]["task_id"]


def _prepare_uploads(tmp_path: Path, monkeypatch) -> Path:
    uploads = tmp_path / "uploads"
    uploads.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(upload, "UPLOAD_DIR", uploads)
    return uploads


def test_detect_start_applies_the_requested_rule_set(tmp_path, monkeypatch) -> None:
    _prepare_uploads(tmp_path, monkeypatch)
    monkeypatch.setenv("PAPER_RULES_DIR", str(_prepare_rules_dir(tmp_path, second=_second_rule_payload())))
    client = TestClient(app)

    task_id = _upload_sample(client)
    started = client.post("/api/v1/detect/start", json={"task_id": task_id, "rule_set_id": SECOND_ID})

    assert started.status_code == 200
    assert started.json()["data"]["rule_set_id"] == SECOND_ID
    errors = client.get(f"/api/v1/detect/result/{task_id}").json()["data"]["errors"]
    # 选中的规则集真的生效：正文是宋体，而 school-b 要求黑体
    assert any(
        error["type"] == "font_error" and error["expected"] == "黑体" and "正文内容示例" in error["content"]
        for error in errors
    )


def test_detect_without_rule_set_uses_the_default(tmp_path, monkeypatch) -> None:
    _prepare_uploads(tmp_path, monkeypatch)
    monkeypatch.setenv("PAPER_RULES_DIR", str(_prepare_rules_dir(tmp_path, second=_second_rule_payload())))
    monkeypatch.delenv("PAPER_RULES_DEFAULT", raising=False)
    client = TestClient(app)

    task_id = _upload_sample(client)
    started = client.post("/api/v1/detect/start", json={"task_id": task_id})

    assert started.json()["data"]["rule_set_id"] == "default"
    errors = client.get(f"/api/v1/detect/result/{task_id}").json()["data"]["errors"]
    # 默认规则要求宋体，样例正文正是宋体 → 不该报字体错
    assert not any(error["type"] == "font_error" and "正文内容示例" in error["content"] for error in errors)


def test_detect_start_rejects_unknown_rule_set(tmp_path, monkeypatch) -> None:
    _prepare_uploads(tmp_path, monkeypatch)
    monkeypatch.setenv("PAPER_RULES_DIR", str(REPO_RULES))
    client = TestClient(app)
    task_id = _upload_sample(client)

    response = client.post("/api/v1/detect/start", json={"task_id": task_id, "rule_set_id": "no-such-school"})

    assert response.status_code == 404
    assert "规则集不存在" in response.json()["detail"]


def test_report_repeats_the_rule_set_chosen_at_detect_time(tmp_path, monkeypatch) -> None:
    """检测选了别的规则集，报告就不能回去复述默认规则集的依据。"""
    _prepare_uploads(tmp_path, monkeypatch)
    monkeypatch.setattr(report_service, "REPORT_DIR", tmp_path / "reports")
    monkeypatch.setenv("PAPER_RULES_DIR", str(_prepare_rules_dir(tmp_path, second=_second_rule_payload())))
    client = TestClient(app)
    task_id = _upload_sample(client)
    client.post("/api/v1/detect/start", json={"task_id": task_id, "rule_set_id": SECOND_ID})
    try:
        response = client.post("/api/v1/report/create", json={"task_id": task_id})
    finally:
        store.delete(task_id)

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["rule_set_id"] == SECOND_ID
    assert data["rule_file"] == f"{SECOND_ID}.json"
