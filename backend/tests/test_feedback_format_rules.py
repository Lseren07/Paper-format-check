from backend.app.models.contracts import Document
from backend.app.parser.structure import heading_level_from_text
from backend.app.parser.word_parser import parse_docx
from backend.app.rules.contracts import CheckRule
from backend.app.rules.detectors import detect_alignment, detect_font, detect_toc_consistency
from backend.app.rules.spec_detectors import detect_keyword_format
from backend.app.rules.targets import paragraph_targets
from io import BytesIO
from docx import Document as BuildDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def _paragraph(*, paragraph_id='p-0001', text='', style='Normal', heading=None, alignment=None, structure='body', indent_chars=None, runs=None, **extra):
    payload = {
        'paragraph_id': paragraph_id,
        'text': text,
        'style': {'name': style},
        'heading': {'level': heading},
        'structure': structure,
        'structure_role': extra.pop('role', None),
        'location': {'part': 'document'},
        'format': {'alignment': alignment},
        'runs': runs or [{'text': text, 'font': {'effective': '宋体', 'ascii': '宋体', 'east_asia': '宋体'}, 'size_pt': 12, 'bold': False}],
    }
    if indent_chars is not None:
        payload['format']['first_line_indent_chars'] = indent_chars
    payload.update(extra)
    return payload


def _document(paragraphs, toc=None):
    return Document(document_id='D1', source_filename='x.docx', paragraphs=paragraphs, metadata={'toc_paragraphs': toc or []})


def _set_run_fonts(run, *, ascii_name: str, east_asia: str) -> None:
    run.font.name = ascii_name
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.get_or_add_rFonts()
    fonts.set(qn('w:ascii'), ascii_name)
    fonts.set(qn('w:hAnsi'), ascii_name)
    fonts.set(qn('w:eastAsia'), east_asia)


def test_keyword_format_requires_flush_left_and_bold_label() -> None:
    document = _document([_paragraph(
        text='关键词：数据库，性能，检测',
        structure='keywords',
        alignment='center',
        indent_chars=2,
        runs=[
            {'text': '关键词：', 'font': {'effective': '宋体', 'ascii': '宋体', 'east_asia': '宋体'}, 'size_pt': 12, 'bold': False},
            {'text': '数据库，性能，检测', 'font': {'effective': '宋体', 'ascii': '宋体', 'east_asia': '宋体'}, 'size_pt': 12, 'bold': False},
        ],
    )])
    rule = CheckRule(id='keywords-zh', type='keyword_format', target='keywords', expected={
        'pattern': r'^关键词\s*[：:]',
        'min': 3,
        'max': 8,
        'required': True,
        'alignment': 'left',
        'first_line_indent': 0,
        'label': {'font': '宋体', 'size': '小四', 'bold': True},
        'content': {'font': '宋体', 'size': '小四'},
    })
    errors = detect_keyword_format(document, rule)
    assert any(error.expected == 'left' and error.current == 'center' for error in errors)
    assert any(error.expected in {'True', 'true', '加粗'} or 'bold' in error.expected.lower() or error.expected == 'True' for error in errors) or any(error.current in {'False', 'false', '未加粗'} for error in errors)


def test_english_keyword_label_uses_times_and_content_stays_times() -> None:
    document = _document([_paragraph(
        text='Keywords: database, performance, testing',
        structure='keywords_en',
        alignment='left',
        runs=[
            {'text': 'Keywords: ', 'font': {'effective': 'Times New Roman', 'ascii': 'Times New Roman', 'east_asia': '宋体'}, 'size_pt': 12, 'bold': True},
            {'text': 'database, performance, testing', 'font': {'effective': 'Times New Roman', 'ascii': 'Times New Roman', 'east_asia': '宋体'}, 'size_pt': 12, 'bold': False},
        ],
    )])
    rule = CheckRule(id='keywords-en-format', type='keyword_format', target='keywords-en', expected={
        'pattern': r'^Keywords\s*[：:]',
        'min': 3,
        'max': 8,
        'alignment': 'left',
        'label': {'font': 'Times New Roman', 'size': '小四', 'bold': True},
        'content': {'font': 'Times New Roman', 'size': '小四'},
    })
    assert detect_keyword_format(document, rule) == []


def test_numbered_second_level_title_is_not_checked_as_body_alignment() -> None:
    paragraph = _paragraph(paragraph_id='p-0046', text='1.1 研究背景', alignment='center', heading=None, style='Normal')
    document = _document([paragraph])
    assert heading_level_from_text('1.1 研究背景') == 2
    assert 'title2' in paragraph_targets(document)['p-0046']
    assert 'body' not in paragraph_targets(document)['p-0046']
    body_rule = CheckRule(id='body-alignment', type='alignment', target='body', expected={'alignment': 'justify'})
    title_rule = CheckRule(id='title2-alignment', type='alignment', target='title2', expected={'alignment': 'left'})
    assert detect_alignment(document, body_rule) == []
    errors = detect_alignment(document, title_rule)
    assert len(errors) == 1
    assert errors[0].current == 'center'
    assert errors[0].expected == 'left'


def test_mixed_sentence_checks_cjk_and_latin_fonts_separately() -> None:
    text = 'PostgreSQL 和 MySQL 数据库的性能基准测试'
    document = _document([_paragraph(
        text=text,
        runs=[{
            'text': text,
            'font': {'effective': '宋体', 'ascii': 'Times New Roman', 'east_asia': '宋体'},
            'size_pt': 12,
            'bold': False,
        }],
    )])
    body_rule = CheckRule(id='body-font', type='font', target='body', expected={'font': '宋体'})
    english_rule = CheckRule(id='foreign-body-font', type='font', target='body', expected={'font': 'Times New Roman'})
    assert detect_font(document, body_rule) == []
    assert detect_font(document, english_rule) == []


def test_mixed_sentence_reports_latin_font_when_ascii_is_songti() -> None:
    text = 'PostgreSQL 和 MySQL 数据库的性能基准测试'
    document = _document([_paragraph(
        text=text,
        runs=[{
            'text': text,
            'font': {'effective': '宋体', 'ascii': '宋体', 'east_asia': '宋体'},
            'size_pt': 12,
            'bold': False,
        }],
    )])
    rule = CheckRule(id='body-font', type='font', target='body', expected={'font': '宋体'})
    errors = detect_font(document, rule)
    assert errors
    assert errors[0].current == '宋体'
    assert errors[0].expected == 'Times New Roman'


def test_toc_consistency_compares_titles_by_heading_number() -> None:
    document = _document(
        [_paragraph(paragraph_id='p-0001', text='1.2 其他内容', heading=2, style='Heading 2')],
        toc=[('1.2国内外研究问题', 2)],
    )
    errors = detect_toc_consistency(document, CheckRule(id='toc-consistency', type='toc_consistency', expected={}))
    assert errors
    assert '国内外研究问题' in (errors[0].expected + errors[0].current + errors[0].content)


def test_toc_consistency_accepts_same_title_with_page_number() -> None:
    document = _document(
        [_paragraph(paragraph_id='p-0001', text='1.2 国内外研究问题', heading=2, style='Heading 2')],
        toc=[('1.2国内外研究问题\t12', 2)],
    )
    assert detect_toc_consistency(document, CheckRule(id='toc-consistency', type='toc_consistency', expected={})) == []


def test_parser_marks_plain_numbered_subtitle_as_heading_two() -> None:
    docx = BuildDocument()
    paragraph = docx.add_paragraph('1.1 研究背景')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run_fonts(paragraph.runs[0], ascii_name='Times New Roman', east_asia='黑体')
    stream = BytesIO()
    docx.save(stream)
    parsed = parse_docx(stream.getvalue(), source_filename='x.docx')
    item = next(p for p in parsed.paragraphs if '1.1' in (p.get('text') or ''))
    assert item['heading']['level'] == 2
