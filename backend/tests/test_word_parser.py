from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile
from docx import Document as BuildDocument
from docx.shared import Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from PIL import Image

import pytest

from backend.app.parser.errors import DocumentParseError
from backend.app.parser.word_parser import parse_docx
from backend.app.parser.numbering import numbering_label
from backend.app.parser.styles import parse_styles, parse_theme, resolve_style
from xml.etree import ElementTree as ET


def make_docx_xml(document_xml: str = "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p><w:r><w:t>Hello</w:t></w:r></w:p></w:body></w:document>") -> bytes:
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'><Default Extension='rels' ContentType='application/vnd.openxmlformats-package.relationships+xml'/><Default Extension='xml' ContentType='application/xml'/><Override PartName='/word/document.xml' ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'/></Types>")
        archive.writestr("_rels/.rels", "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'><Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument' Target='word/document.xml'/></Relationships>")
        archive.writestr("word/document.xml", document_xml)
    return stream.getvalue()


def test_parse_docx_extracts_paragraph_text_and_run() -> None:
    document = parse_docx(BytesIO(make_docx_xml()), source_filename="thesis.docx", document_id="DOC001")
    assert document.document_id == "DOC001"
    assert document.source_filename == "thesis.docx"
    assert document.paragraphs[0]["text"] == "Hello"
    assert document.paragraphs[0]["runs"][0]["text"] == "Hello"


def test_parse_docx_rejects_missing_document_xml() -> None:
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
    with pytest.raises(DocumentParseError):
        parse_docx(BytesIO(stream.getvalue()))


def test_parse_docx_extracts_headers_and_footers() -> None:
    source = BuildDocument()
    source.add_paragraph("正文")
    source.sections[0].header.paragraphs[0].text = "页眉"
    source.sections[0].footer.paragraphs[0].text = "页脚"
    stream = BytesIO(); source.save(stream)
    parsed = parse_docx(BytesIO(stream.getvalue()))
    assert parsed.headers[0]["text"] == "页眉"
    assert parsed.footers[0]["text"] == "页脚"


def test_parse_numbered_heading_exposes_numbering_metadata() -> None:
    source = BuildDocument()
    p = source.add_paragraph("第一节")
    p.style = "Heading 1"
    stream = BytesIO(); source.save(stream)
    parsed = parse_docx(BytesIO(stream.getvalue()))
    assert parsed.paragraphs[0]["heading"]["level"] == 1


def test_numbering_label_expands_decimal_multilevel_template() -> None:
    assert numbering_label("%1.%2", [2, 3], "decimal") == "2.3"
    assert numbering_label("%1", [4], "upperRoman") == "IV"
    assert numbering_label("%1", [3], "lowerLetter") == "c"


def test_styles_resolve_doc_defaults_and_theme_font() -> None:
    xml = """<w:styles xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:eastAsia='宋体'/><w:sz w:val='24'/></w:rPr></w:rPrDefault></w:docDefaults><w:style w:styleId='Body'><w:name w:val='正文'/><w:rPr><w:rFonts w:ascii='Arial'/></w:rPr></w:style></w:styles>"""
    parsed = parse_styles(ET.fromstring(xml))
    assert parsed["defaults"]["font"]["eastAsia"] == "宋体"
    assert resolve_style("Body", parsed)["font"]["ascii"] == "Arial"


def test_theme_font_resolves_minor_east_asia_reference() -> None:
    theme = parse_theme(ET.fromstring("""<a:theme xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'><a:themeElements><a:fontScheme><a:minorFont><a:latin typeface='Calibri'/><a:ea typeface='等线'/></a:minorFont></a:fontScheme></a:themeElements></a:theme>"""))
    styles = {"styles": {"Body": {"based_on": None, "rpr": {"font_theme": {"eastAsia": "minorEastAsia"}}}}, "defaults": {}, "theme": theme}
    assert resolve_style("Body", styles)["font"]["eastAsia"] == "等线"


def test_toc_style_is_recorded_in_metadata() -> None:
    source = BuildDocument(); source.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH); p = source.add_paragraph("第一章 绪论"); p.style = "TOC 1"
    stream = BytesIO(); source.save(stream)
    parsed = parse_docx(BytesIO(stream.getvalue()))
    assert parsed.metadata["toc_paragraphs"][0]["text"] == "第一章 绪论"


def test_parse_docx_keeps_drawing_relationship_metadata() -> None:
    source = BuildDocument(); source.add_paragraph("图示")
    image = BytesIO(); Image.new("RGB", (20, 10), "red").save(image, format="PNG"); image.seek(0)
    source.paragraphs[0].add_run().add_picture(image, width=Inches(1))
    stream = BytesIO(); source.save(stream)
    parsed = parse_docx(BytesIO(stream.getvalue()))
    drawing = parsed.paragraphs[0]["drawings"][0]
    assert drawing["target"].endswith(".png")
    assert drawing["width_emu"] > 0
    assert drawing["height_emu"] > 0


def test_parse_docx_extracts_page_margins_and_table_cell_paragraphs() -> None:
    source = BuildDocument()
    source.sections[0].top_margin = Inches(1)
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "表格正文"
    cell.paragraphs[0].runs[0].font.name = "黑体"
    ind = cell.paragraphs[0]._p.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), "200")
    stream = BytesIO(); source.save(stream)
    parsed = parse_docx(BytesIO(stream.getvalue()))
    table_paragraph = next(paragraph for paragraph in parsed.paragraphs if paragraph["location"]["part"] == "table")
    assert parsed.sections[0]["margins_pt"]["top"] == 72
    assert parsed.sections[0]["header_distance_pt"] is not None
    assert parsed.sections[0]["footer_distance_pt"] is not None
    assert table_paragraph["text"] == "表格正文"
    assert table_paragraph["format"]["first_line_indent_chars"] == 2

def _save_docx(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _add_page_field(paragraph) -> None:
    from docx.oxml import OxmlElement

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    paragraph.add_run()._r.append(instr)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def test_parse_docx_identifies_cover_abstract_and_references() -> None:
    source = BuildDocument()
    source.add_paragraph("电子科技大学成都学院")
    title = source.add_paragraph("面向格式检测的论文题目")
    title.runs[0].font.size = Inches(0.22)
    source.add_paragraph("摘要")
    source.add_paragraph("这是摘要正文，用于结构识别。")
    source.add_paragraph("关键词：格式；检测")
    heading = source.add_paragraph("第一章 绪论")
    heading.style = "Heading 1"
    source.add_paragraph("这是正文段落。")
    source.add_paragraph("图1 系统架构")
    source.add_paragraph("参考文献")
    source.add_paragraph("[1] 张三. 测试文献.")
    parsed = parse_docx(BytesIO(_save_docx(source)))
    by_text = {item["text"]: item for item in parsed.paragraphs if item["location"]["part"] == "document"}
    assert by_text["电子科技大学成都学院"]["structure"] == "cover"
    assert by_text["面向格式检测的论文题目"]["structure"] == "cover"
    assert by_text["面向格式检测的论文题目"]["structure_role"] == "title"
    assert by_text["摘要"]["structure"] == "abstract"
    assert by_text["摘要"]["structure_role"] == "title"
    assert by_text["这是摘要正文，用于结构识别。"]["structure"] == "abstract"
    assert by_text["关键词：格式；检测"]["structure"] == "keywords"
    assert by_text["这是正文段落。"]["structure"] == "body"
    assert by_text["图1 系统架构"]["structure"] == "figure_caption"
    assert by_text["[1] 张三. 测试文献."]["structure"] == "references"
    structure = parsed.metadata["structure"]
    assert "面向格式检测的论文题目" in structure["cover"]["title"]
    assert structure["abstract"]["paragraph_ids"]
    assert structure["references"]["paragraph_ids"]


def test_parse_docx_does_not_invent_cover_without_abstract_marker() -> None:
    source = BuildDocument()
    source.add_paragraph("第二章 系统设计", style="Heading 1")
    source.add_paragraph("随着人工智能技术的发展")
    parsed = parse_docx(BytesIO(_save_docx(source)))
    assert all(item.get("structure") != "cover" for item in parsed.paragraphs)


def test_parse_docx_extracts_page_number_fields() -> None:
    from docx.oxml import OxmlElement

    source = BuildDocument()
    source.add_paragraph("正文")
    footer = source.sections[0].footer.paragraphs[0]
    footer.alignment = 1
    _add_page_field(footer)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), "upperRoman")
    source.sections[0]._sectPr.append(pg_num)
    parsed = parse_docx(BytesIO(_save_docx(source)))
    assert parsed.pages
    page = parsed.pages[0]
    assert page["section_index"] == 0
    assert page["number_format"] == "upperRoman"
    assert page["position"] == "footer"
    assert any("PAGE" in str(field.get("instruction", "")).upper() for field in page["fields"])


def test_parse_docx_records_table_caption_block_order() -> None:
    source = BuildDocument()
    source.add_paragraph("表1 对比结果")
    source.add_table(rows=1, cols=2)
    parsed = parse_docx(BytesIO(_save_docx(source)))
    caption = next(item for item in parsed.paragraphs if item["text"] == "表1 对比结果")
    assert caption["structure"] == "table_caption"
    assert parsed.tables[0]["block_index"] == caption["block_index"] + 1


def _add_last_rendered_page_break(paragraph) -> None:
    from docx.oxml import OxmlElement
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run._r.insert(0, OxmlElement('w:lastRenderedPageBreak'))


def _add_seq_field(paragraph, name: str) -> None:
    from docx.oxml import OxmlElement
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    run._r.append(begin)
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = f' SEQ {name} \\* ARABIC '
    paragraph.add_run()._r.append(instr)
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    paragraph.add_run()._r.append(end)


def test_parse_docx_assigns_word_rendered_page_numbers() -> None:
    from docx.oxml import OxmlElement
    source = BuildDocument()
    source.add_paragraph('第一页')
    second = source.add_paragraph('第二页')
    _add_last_rendered_page_break(second)
    pg_num = OxmlElement('w:pgNumType')
    pg_num.set(qn('w:start'), '3')
    source.sections[0]._sectPr.append(pg_num)
    parsed = parse_docx(BytesIO(_save_docx(source)))
    by_text = {item['text']: item for item in parsed.paragraphs}
    assert by_text['第一页']['page_number'] == 3
    assert by_text['第二页']['page_number'] == 4
    assert [page['page_number'] for page in parsed.pages] == [3, 4]
    assert by_text['第二页']['paragraph_id'] in parsed.pages[1]['paragraph_ids']
    assert parsed.pages[0]['source'] == 'last_rendered'


def test_parse_docx_identifies_cover_from_thesis_hints_without_abstract() -> None:
    source = BuildDocument()
    source.add_paragraph('电子科技大学成都学院')
    source.add_paragraph('本科毕业论文')
    source.add_paragraph('学号：20230001')
    heading = source.add_paragraph('第一章 绪论')
    heading.style = 'Heading 1'
    source.add_paragraph('正文开始')
    parsed = parse_docx(BytesIO(_save_docx(source)))
    by_text = {item['text']: item for item in parsed.paragraphs if item['location']['part'] == 'document'}
    assert by_text['本科毕业论文']['structure'] == 'cover'
    assert by_text['学号：20230001']['structure'] == 'cover'
    assert by_text['正文开始']['structure'] == 'body'


def test_parse_docx_identifies_abstract_variants_and_complex_captions() -> None:
    source = BuildDocument()
    source.add_paragraph('中文摘要')
    source.add_paragraph('摘要正文')
    source.add_paragraph('图1-1 系统架构')
    source.add_paragraph('Figure 2 Overview')
    source.add_paragraph('表1.1 对比结果')
    caption = source.add_paragraph('实验结果曲线')
    caption.style = 'Caption'
    _add_seq_field(caption, '图')
    parsed = parse_docx(BytesIO(_save_docx(source)))
    by_text = {item['text']: item for item in parsed.paragraphs if item['location']['part'] == 'document'}
    assert by_text['中文摘要']['structure'] == 'abstract'
    assert by_text['摘要正文']['structure'] == 'abstract'
    assert by_text['图1-1 系统架构']['structure'] == 'figure_caption'
    assert by_text['Figure 2 Overview']['structure'] == 'figure_caption'
    assert by_text['表1.1 对比结果']['structure'] == 'table_caption'
    caption_para = next(item for item in parsed.paragraphs if '实验结果曲线' in item['text'])
    assert caption_para['structure'] == 'figure_caption'



def _wrap_in_sdt(paragraph) -> None:
    from docx.oxml import OxmlElement

    sdt = OxmlElement("w:sdt")
    sdt.append(OxmlElement("w:sdtPr"))
    content = OxmlElement("w:sdtContent")
    element = paragraph._element
    parent = element.getparent()
    parent.replace(element, sdt)
    content.append(element)
    sdt.append(content)


def test_parse_docx_reads_page_field_inside_footer_sdt() -> None:
    source = BuildDocument()
    source.add_paragraph("正文")
    footer = source.sections[0].footer.paragraphs[0]
    _add_page_field(footer)
    _wrap_in_sdt(footer)
    parsed = parse_docx(BytesIO(_save_docx(source)))
    assert parsed.pages
    assert parsed.pages[0]["position"] == "footer"
    assert any("PAGE" in str(field.get("instruction", "")).upper() for field in parsed.pages[0]["fields"])


def test_parse_docx_cover_title_skips_larger_school_name() -> None:
    source = BuildDocument()
    school = source.add_paragraph("电子科技大学成都学院")
    school.runs[0].font.size = Inches(0.5)
    source.add_paragraph("毕业论文（设计）")
    source.add_paragraph("题    目 面向格式检测的论文题目")
    source.add_paragraph("摘要")
    source.add_paragraph("这是摘要正文，用于结构识别。")
    parsed = parse_docx(BytesIO(_save_docx(source)))
    by_text = {item["text"]: item for item in parsed.paragraphs if item["location"]["part"] == "document"}
    assert by_text["电子科技大学成都学院"].get("structure_role") != "title"
    assert by_text["题    目 面向格式检测的论文题目"]["structure_role"] == "title"
    assert "面向格式检测的论文题目" in parsed.metadata["structure"]["cover"]["title"]
