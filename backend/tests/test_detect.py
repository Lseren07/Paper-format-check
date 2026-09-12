from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from backend.app.api import upload
from backend.app.main import app


def make_mismatched_docx() -> bytes:
    document = Document()
    document.add_paragraph("第二章 系统设计", style="Heading 1")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("随着人工智能技术的发展")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def upload_sample(client: TestClient, content: bytes, name: str = "thesis.docx") -> str:
    response = client.post(
        "/api/v1/paper/upload",
        files={"file": (name, content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    return response.json()["data"]["task_id"]


def test_detect_start_unknown_task_returns_404(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(upload, "UPLOAD_DIR", tmp_path)
    response = TestClient(app).post("/api/v1/detect/start", json={"task_id": "Tmissing"})
    assert response.status_code == 404
    assert response.json()["detail"] == "检测任务不存在"


def test_detect_result_unknown_task_returns_404(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(upload, "UPLOAD_DIR", tmp_path)
    response = TestClient(app).get("/api/v1/detect/result/Tmissing")
    assert response.status_code == 404
    assert response.json()["detail"] == "检测任务不存在"


def test_detect_result_before_start_returns_409(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(upload, "UPLOAD_DIR", tmp_path)
    client = TestClient(app)
    task_id = upload_sample(client, make_mismatched_docx())
    response = client.get(f"/api/v1/detect/result/{task_id}")
    assert response.status_code == 409
    assert response.json()["detail"] == "检测尚未完成"


def test_upload_then_detect_returns_font_errors(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(upload, "UPLOAD_DIR", tmp_path)
    client = TestClient(app)
    task_id = upload_sample(client, make_mismatched_docx(), "thesis.docx")

    started = client.post("/api/v1/detect/start", json={"task_id": task_id})
    assert started.status_code == 200
    body = started.json()
    assert body["code"] == 200
    assert body["data"]["task_id"] == task_id
    assert body["data"]["status"] == "completed"

    status = client.get(f"/api/v1/detect/status/{task_id}")
    assert status.status_code == 200
    assert status.json()["data"]["status"] == "completed"
    assert status.json()["data"]["progress"] == 100

    result = client.get(f"/api/v1/detect/result/{task_id}")
    assert result.status_code == 200
    data = result.json()["data"]
    assert data["task_id"] == task_id
    assert data["filename"] == "thesis.docx"
    assert data["total_error"] == len(data["errors"])
    assert data["total_error"] >= 1
    font_error = next(
        error
        for error in data["errors"]
        if error["type"] == "font_error" and "随着人工智能技术的发展" in error["content"]
    )
    assert font_error["current"] == "黑体"
    assert font_error["expected"] == "宋体"
    assert "随着人工智能技术的发展" in font_error["content"]
    assert font_error["location"] == "第二章 系统设计"
    assert font_error["error_id"]

