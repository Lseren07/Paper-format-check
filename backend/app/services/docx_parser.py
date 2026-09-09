"""Extract format-relevant structure from a DOCX document."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


_HEADING_RE = re.compile(r"^(?:heading|标题)\s*([1-9])$", re.IGNORECASE)
_ALIGNMENTS = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    WD_ALIGN_PARAGRAPH.DISTRIBUTE: "distribute",
}
_LINE_SPACING_RULES = {
    WD_LINE_SPACING.SINGLE: "single",
    WD_LINE_SPACING.ONE_POINT_FIVE: "1.5",
    WD_LINE_SPACING.DOUBLE: "double",
    WD_LINE_SPACING.AT_LEAST: "at_least",
    WD_LINE_SPACING.EXACTLY: "exactly",
    WD_LINE_SPACING.MULTIPLE: "multiple",
}


def parse_docx(path: str | Path) -> dict[str, Any]:
    """Return a JSON-serializable representation of a DOCX document."""

    source = Path(path)
    document = Document(source)
    return {
        "schema_version": "1.0",
        "source_filename": source.name,
        "sections": [_parse_section(index, section) for index, section in enumerate(document.sections)],
        "paragraphs": [
            _parse_paragraph(index, paragraph)
            for index, paragraph in enumerate(document.paragraphs)
        ],
        "tables": [_parse_table(index, table) for index, table in enumerate(document.tables)],
        "headers": _parse_headers(document),
        "footers": _parse_footers(document),
    }


def _parse_section(index: int, section: Any) -> dict[str, Any]:
    return {
        "index": index,
        "page_width_cm": _length_cm(section.page_width),
        "page_height_cm": _length_cm(section.page_height),
        "margins_cm": {
            "top": _length_cm(section.top_margin),
            "right": _length_cm(section.right_margin),
            "bottom": _length_cm(section.bottom_margin),
            "left": _length_cm(section.left_margin),
        },
        "header_distance_cm": _length_cm(section.header_distance),
        "footer_distance_cm": _length_cm(section.footer_distance),
    }


def _parse_paragraph(index: int, paragraph: Any) -> dict[str, Any]:
    formatting = paragraph.paragraph_format
    return {
        "index": index,
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style else None,
        "heading_level": _heading_level(paragraph),
        "alignment": _paragraph_value(paragraph, "alignment", _ALIGNMENTS),
        "space_before_pt": _paragraph_value(paragraph, "space_before", _points),
        "space_after_pt": _paragraph_value(paragraph, "space_after", _points),
        "line_spacing_pt": _line_spacing_points(paragraph),
        "line_spacing_rule": _paragraph_value(paragraph, "line_spacing_rule", _LINE_SPACING_RULES),
        "first_line_indent_pt": _paragraph_value(paragraph, "first_line_indent", _points),
        "left_indent_pt": _paragraph_value(paragraph, "left_indent", _points),
        "right_indent_pt": _paragraph_value(paragraph, "right_indent", _points),
        "runs": [_parse_run(run) for run in paragraph.runs],
    }


def _parse_run(run: Any) -> dict[str, Any]:
    return {
        "text": run.text,
        "font": _run_font_name(run),
        "size_pt": _run_value(run, "size", _points),
        "bold": _run_value(run, "bold", lambda value: value),
        "italic": _run_value(run, "italic", lambda value: value),
        "underline": _run_value(run, "underline", lambda value: value),
    }


def _parse_table(index: int, table: Any) -> dict[str, Any]:
    return {
        "index": index,
        "rows": [[cell.text for cell in row.cells] for row in table.rows],
    }


def _parse_headers(document: Any) -> list[dict[str, Any]]:
    return [
        {"section_index": index, "text": section.header.paragraphs[0].text if section.header.paragraphs else ""}
        for index, section in enumerate(document.sections)
    ]


def _parse_footers(document: Any) -> list[dict[str, Any]]:
    return [
        {"section_index": index, "text": section.footer.paragraphs[0].text if section.footer.paragraphs else ""}
        for index, section in enumerate(document.sections)
    ]


def _heading_level(paragraph: Any) -> int | None:
    style_name = paragraph.style.name if paragraph.style else ""
    match = _HEADING_RE.match(style_name.strip())
    if match:
        return int(match.group(1))

    outline_level = paragraph._p.pPr.find(qn("w:outlineLvl")) if paragraph._p.pPr is not None else None
    if outline_level is not None:
        value = outline_level.get(qn("w:val"))
        if value and value.isdigit():
            return int(value) + 1
    return None


def _font_name(run: Any) -> str | None:
    r_pr = run._element.rPr
    if r_pr is None:
        return run.font.name
    fonts = r_pr.rFonts
    if fonts is None:
        return run.font.name
    for attribute in ("eastAsia", "hAnsi", "ascii", "cs"):
        value = fonts.get(qn(f"w:{attribute}"))
        if value:
            return value
    return run.font.name


def _length_cm(value: Any) -> float | None:
    return round(float(value.cm), 2) if value is not None else None


def _points(value: Any) -> float | None:
    return round(float(value.pt), 2) if value is not None else None


def _line_spacing_points(paragraph: Any) -> float | None:
    value = _paragraph_value(paragraph, "line_spacing", lambda item: item)
    if value is None or isinstance(value, float):
        return None
    return _points(value)


def _paragraph_value(paragraph: Any, attribute: str, transform: Any) -> Any:
    value = getattr(paragraph.paragraph_format, attribute)
    if value is not None:
        return _transform(value, transform)

    style = paragraph.style
    while style is not None:
        value = getattr(style.paragraph_format, attribute)
        if value is not None:
            return _transform(value, transform)
        style = style.base_style
    return None


def _run_value(run: Any, attribute: str, transform: Any) -> Any:
    value = getattr(run.font, attribute)
    if value is not None:
        return _transform(value, transform)

    style = run.style
    while style is not None:
        value = getattr(style.font, attribute)
        if value is not None:
            return _transform(value, transform)
        style = style.base_style
    return None


def _run_font_name(run: Any) -> str | None:
    direct = _font_name(run)
    if direct is not None:
        return direct

    style = run.style
    while style is not None:
        value = _font_name_from_style(style)
        if value is not None:
            return value
        style = style.base_style
    return None


def _transform(value: Any, transform: Any) -> Any:
    if isinstance(transform, dict):
        return transform.get(value)
    return transform(value)


def _font_name_from_style(style: Any) -> str | None:
    r_pr = style._element.rPr
    if r_pr is None or r_pr.rFonts is None:
        return style.font.name
    fonts = r_pr.rFonts
    for attribute in ("eastAsia", "hAnsi", "ascii", "cs"):
        value = fonts.get(qn(f"w:{attribute}"))
        if value:
            return value
    return style.font.name
